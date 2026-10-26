"""导入中心服务（SRS-D01.06 FR-01/FR-02/FR-03）

执行流水线（§5.1）：
  流式解析 → 框架层校验（零业务代码）→ 业务校验（导入器 rowValidate，逐行独立）
  → 批次提交（批=事务边界，批内原子/批间独立）→ 汇总与错误报告

MVP 简化（相对 SRS 的差异，均已标注）：
- 文件存本地目录（backend/var/imports/）替代 OSS 分片上传；source_file_ref 存相对路径
- 任务在请求内同步执行，替代 D01.01 worker 异步认领（断点续跑能力已具备）
- 模板以代码注册为权威，登记到 oms_base.bas_ie_template（与数据权限注册表同策略）
"""
import json
import pathlib
import time
from datetime import datetime, timedelta
from typing import Any

from redis.asyncio import Redis
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import ConflictError, NotFoundError, ValidationError
from app.core.importer import (
    ImporterInterface,
    RowError,
    TemplateSpec,
    build_template_file,
    read_rows,
    rows_to_xlsx,
    validate_row_format,
)
from app.models.base import BaseImportError, BaseImportTask

BATCH_SIZE = 1000
# ⚠ 层级：app/modules/m2_stock/service/x.py → parents[4] 才是 backend/
#   （P1 目录重组后 service 下沉了两层；parents[2] 会指向 modules/ → 产物落到
#    app/modules/var/ 里，既是垃圾目录又不在 .gitignore 覆盖范围）
#   运行时产物统一放 backend/var/（.gitignore 已覆盖 backend/var/）
UPLOAD_DIR = pathlib.Path(__file__).resolve().parents[4] / "var" / "imports"

# ---------- 模板注册表（代码=权威，DB 登记供任务中心筛选） ----------

TEMPLATE_REGISTRY: dict[str, TemplateSpec] = {}


def register_template(spec: TemplateSpec) -> None:
    TEMPLATE_REGISTRY[spec.template_code] = spec


def get_template(code: str) -> TemplateSpec:
    spec = TEMPLATE_REGISTRY.get(code)
    if not spec:
        raise NotFoundError(f"导入模板不存在: {code}")
    return spec


# ---------- 导入器注册表 ----------

IMPORTER_REGISTRY: dict[str, ImporterInterface] = {}


def register_importer(importer: ImporterInterface) -> None:
    IMPORTER_REGISTRY[importer.template_code] = importer


def get_importer(code: str) -> ImporterInterface:
    imp = IMPORTER_REGISTRY.get(code)
    if not imp:
        raise NotFoundError(f"未注册的导入器: {code}")
    return imp


class ImportContext:
    """导入上下文（批量预取缓存 + 任务信息）"""

    def __init__(self, task_id: int, template: TemplateSpec, session: AsyncSession):
        self.task_id = task_id
        self.template = template
        self.session = session
        self.cache: dict[str, Any] = {}


class ImportService:
    """导入任务编排"""

    def __init__(self, session: AsyncSession, redis: Redis):
        self.s = session
        self.r = redis

    # ---------- 模板 ----------

    async def list_templates(self) -> list[TemplateSpec]:
        return list(TEMPLATE_REGISTRY.values())

    def template_file(self, code: str) -> bytes:
        return build_template_file(get_template(code))

    # ---------- 任务创建 ----------

    async def _next_task_bn(self) -> str:
        seq = await self.r.incr("seq:import_task")
        return f"IM{datetime.now().strftime('%Y%m%d%H%M%S')}{seq:06d}"

    async def create_task(
        self,
        template_code: str,
        filename: str,
        content: bytes,
        creator: str = "",
        owner_id: int = 0,
        tenant_id: int = 1,
    ) -> BaseImportTask:
        template = get_template(template_code)
        if len(content) > 200 * 1024 * 1024:
            raise ValidationError("文件超过 200MB 上限，请拆分后导入")
        import hashlib

        md5 = hashlib.md5(content).hexdigest()

        # 同 MD5 24h 内重复创建任务 → 提示（FR-03d）
        recent = await self.s.scalar(
            select(BaseImportTask.id).where(
                BaseImportTask.file_md5 == md5,
                BaseImportTask.create_time >= _days_ago(1),
            )
        )
        if recent:
            raise ConflictError(
                "相同内容的文件在 24 小时内已创建过导入任务（可能是重复上传）",
                data={"existing_task_id": recent},
            )

        task_bn = await self._next_task_bn()
        UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        ext = pathlib.Path(filename).suffix or ".xlsx"
        rel_path = f"{task_bn}{ext}"
        (UPLOAD_DIR / rel_path).write_bytes(content)

        task = BaseImportTask(
            tenant_id=tenant_id,
            owner_id=owner_id,
            task_bn=task_bn,
            template_code=template.template_code,
            template_version=template.version,
            entity_ref=template.entity_ref,
            source_file_ref=rel_path,
            file_md5=md5,
            status=0,
            creator=creator,
        )
        self.s.add(task)
        await self.s.commit()
        await self.s.refresh(task)
        return task

    # ---------- 执行 ----------

    async def run_task(self, task_id: int) -> BaseImportTask:
        """执行导入任务（同步执行；支持断点续跑）"""
        task = await self._get_task(task_id)
        if task.status == 1:
            raise ConflictError("任务正在执行中")
        if task.status == 2:
            raise ConflictError("任务已全部成功，无需重复执行")

        template = get_template(task.template_code)
        importer = get_importer(task.template_code)
        path = UPLOAD_DIR / (task.source_file_ref or "")
        if not path.exists():
            raise NotFoundError(f"源文件不存在: {task.source_file_ref}")

        started = time.time()
        task.status = 1
        task.start_time = _now()
        await self.s.commit()

        try:
            headers, data_rows = read_rows(path.read_bytes(), path.name)
        except Exception as exc:  # 解析失败 = 系统级失败（status=4）
            task.status = 4
            task.fail_reason = f"文件解析失败: {exc}"[:512]
            task.end_time = _now()
            task.cost_ms = int((time.time() - started) * 1000)
            await self.s.commit()
            await self.s.refresh(task)
            return task

        # 表头 → 字段映射
        header_map = _map_headers(headers, template)
        task.row_count_total = len(data_rows)

        ctx = ImportContext(task.id, template, self.s)
        checkpoint = int((task.progress_json or {}).get("checkpoint_batch", 0))
        processed = 0
        success_rows = 0
        fail_rows = 0
        error_records: list[tuple[int, str, str, str, Any]] = []
        batch_no = 0

        for start in range(0, len(data_rows), BATCH_SIZE):
            batch_no += 1
            if batch_no <= checkpoint:
                processed += len(data_rows[start:start + BATCH_SIZE])
                continue
            batch = data_rows[start:start + BATCH_SIZE]
            row_dicts: list[dict[str, Any]] = []
            for offset, raw in enumerate(batch):
                row_no = start + offset + 2  # 表头占第 1 行
                row_dicts.append(
                    {
                        "_row_no": row_no,
                        "_raw": raw,
                        **{c.field_map: _pick(raw, header_map, c.field_map) for c in template.columns},
                    }
                )

            # ① 框架层校验（零业务代码）→ 通过者进入业务校验候选
            candidates: list[dict] = []
            for row in row_dicts:
                row_no = row["_row_no"]
                fmt = validate_row_format(row, template, row_no)
                if not fmt.ok:
                    fail_rows += 1
                    error_records.append((row_no, "format", fmt.code, fmt.message, row["_raw"]))
                    continue
                candidates.append(row)

            # ② 批量预取（性能契约：杜绝 N+1 查库）——必须在业务校验之前，
            #    prefetch 负责把本批引用键一次性装进 ctx.cache
            await importer.prefetch(candidates, ctx)

            # ③ 业务层校验（逐行独立，单行错误不阻塞其他行）
            pass_rows: list[dict] = []
            for row in candidates:
                row_no = row["_row_no"]
                try:
                    res = importer.row_validate(row, ctx)
                    if not res.ok:
                        fail_rows += 1
                        error_records.append((row_no, "business", res.code, res.message, row["_raw"]))
                        continue
                except RowError as exc:
                    fail_rows += 1
                    error_records.append((row_no, "business", exc.code, exc.message, row["_raw"]))
                    continue
                pass_rows.append(row)

            # 批次提交（批=事务边界）
            if pass_rows:
                try:
                    await importer.commit_batch(pass_rows, ctx)
                    await self.s.commit()
                    success_rows += len(pass_rows)
                except Exception as exc:
                    await self.s.rollback()
                    fail_rows += len(pass_rows)
                    for row in pass_rows:
                        error_records.append(
                            (
                                row["_row_no"], "commit", "COMMIT_FAILED",
                                f"第{row['_row_no']}行提交失败：{exc}"[:512], row["_raw"],
                            )
                        )

            processed += len(batch)
            # 进度上报（批次粒度）
            task.progress_json = {
                "processed": processed,
                "success": success_rows,
                "fail": fail_rows,
                "percent": round(processed * 100 / max(len(data_rows), 1), 2),
                "checkpoint_batch": batch_no,
                "est_remain": 0,
            }
            await self.s.commit()

        # 错误明细落库
        for row_no, stage, code, message, raw in error_records:
            self.s.add(
                BaseImportError(
                    tenant_id=task.tenant_id,
                    task_id=task.id,
                    row_no=row_no,
                    error_stage=stage,
                    error_code=code,
                    error_message=message[:512],
                    row_snapshot=_jsonable(raw),
                )
            )
        await self.s.commit()

        # 错误报告（三 sheet）
        report_ref = ""
        if error_records:
            report_ref = await self._build_error_report(task, error_records, processed, success_rows)

        # 错误聚类摘要
        clusters: dict[str, dict[str, Any]] = {}  # 错误聚类摘要
        for _row_no, _stage, code, _msg, _raw in error_records:
            c = clusters.setdefault(code, {"error_code": code, "count": 0, "sample_rows": []})
            c["count"] += 1
            if len(c["sample_rows"]) < 10:
                c["sample_rows"].append(_row_no)

        task.result_json = {
            "success_rows": success_rows,
            "fail_rows": fail_rows,
            "error_report_ref": report_ref,
            "cluster_summary": sorted(clusters.values(), key=lambda x: -x["count"])[:20],
        }
        task.error_report_ref = report_ref
        task.status = 2 if fail_rows == 0 else (3 if success_rows > 0 else 4)
        task.end_time = _now()
        task.cost_ms = int((time.time() - started) * 1000)
        await self.s.commit()
        await self.s.refresh(task)
        return task

    async def _build_error_report(
        self,
        task: BaseImportTask,
        error_records: list[tuple[int, str, str, str, Any]],
        processed: int,
        success_rows: int,
    ) -> str:
        """生成错误报告 Excel（sheet1 错误明细 / sheet2 聚类摘要 / sheet3 成功统计）"""
        from openpyxl import Workbook

        wb = Workbook()
        ws1 = wb.active
        ws1.title = "错误明细"
        ws1.append(["行号", "错误阶段", "错误码", "错误消息", "原始数据"])
        clusters: dict[str, int] = {}
        for row_no, stage, code, message, raw in error_records:
            ws1.append([row_no, stage, code, message, json.dumps(_jsonable(raw), ensure_ascii=False)])
            clusters[code] = clusters.get(code, 0) + 1

        ws2 = wb.create_sheet("聚类摘要")
        ws2.append(["错误码", "计数", "说明"])
        for code, cnt in sorted(clusters.items(), key=lambda x: -x[1]):
            ws2.append([code, cnt, _ERROR_CODE_DESC.get(code, "")])

        ws3 = wb.create_sheet("成功统计")
        ws3.append(["总行数", "成功", "失败"])
        ws3.append([processed, success_rows, processed - success_rows])

        UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        ref = f"{task.task_bn}_errors.xlsx"
        wb.save(UPLOAD_DIR / ref)
        wb.close()
        return ref

    # ---------- 查询 ----------

    async def _get_task(self, task_id: int) -> BaseImportTask:
        task = await self.s.scalar(
            select(BaseImportTask).where(BaseImportTask.id == task_id)
        )
        if not task:
            raise NotFoundError(f"导入任务不存在: {task_id}")
        return task

    async def get_task(self, task_id: int) -> BaseImportTask:
        return await self._get_task(task_id)

    async def list_tasks(self, limit: int = 50) -> list[BaseImportTask]:
        rows = await self.s.scalars(
            select(BaseImportTask).order_by(BaseImportTask.id.desc()).limit(limit)
        )
        return list(rows)

    async def list_errors(self, task_id: int) -> list[BaseImportError]:
        rows = await self.s.scalars(
            select(BaseImportError)
            .where(BaseImportError.task_id == task_id)
            .order_by(BaseImportError.row_no)
        )
        return list(rows)


_ERROR_CODE_DESC = {
    "FORMAT_REQUIRED": "必填项为空",
    "FORMAT_TYPE": "类型不匹配",
    "FORMAT_LENGTH": "长度超限",
    "FORMAT_ENUM": "取值不在允许范围",
    "COMMIT_FAILED": "批次提交失败",
}


def _map_headers(headers: list[str], template: TemplateSpec) -> dict[str, int]:
    """表头 → 列索引映射（按中文列名精确匹配，忽略空格）"""
    normalized = {str(h).replace(" ", "").replace("　", ""): i for i, h in enumerate(headers)}
    out: dict[str, int] = {}
    for col in template.columns:
        key = col.col_name.replace(" ", "").replace("　", "")
        if key in normalized:
            out[col.field_map] = normalized[key]
    return out


def _pick(raw: list[Any], header_map: dict[str, int], field_map: str) -> Any:
    idx = header_map.get(field_map)
    if idx is None or idx >= len(raw):
        return ""
    return raw[idx]


def _jsonable(value: Any) -> Any:
    if isinstance(value, (list, tuple)):
        return [_jsonable(v) for v in value]
    if isinstance(value, (datetime,)):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if value is None or isinstance(value, (int, float, str, bool)):
        return value
    return str(value)


def _now() -> datetime:
    return datetime.now()


def _days_ago(n: int) -> datetime:
    return datetime.now() - timedelta(days=n)


__all__ = [
    "BATCH_SIZE",
    "IMPORTER_REGISTRY",
    "TEMPLATE_REGISTRY",
    "ImportContext",
    "ImportService",
    "get_importer",
    "get_template",
    "register_importer",
    "register_template",
    "rows_to_xlsx",
]
