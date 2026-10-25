"""导入中心测试（M1b-2，SRS-D01.06 FR-01/02/03 + D04.01 FR-02d + D04.04 §5.4）

覆盖：模板列表/模板下载/商品导入成功/部分成功（行级错误不阻塞）/货号重复转行错误/
文件内重复/期初库存导入生成单据+记账行+流水/仓未配置可服务货主拒导/同MD5重复创建提示。
"""
import io

import pytest
from sqlalchemy import text

from app.core.db import engine

pytestmark = pytest.mark.asyncio

CLEANUP_SQL = [
    "DELETE FROM oms_base.bas_import_error WHERE task_id IN "
    "(SELECT id FROM oms_base.bas_import_task WHERE task_bn LIKE 'IM%')",
    "DELETE FROM oms_base.bas_import_task WHERE task_bn LIKE 'IM%'",
    "DELETE FROM oms_stock.inv_stock_change_log WHERE bill_bn LIKE 'opening_in%'",
    "DELETE FROM oms_stock.inv_io_bill_item WHERE bill_bn LIKE 'opening_in%'",
    "DELETE FROM oms_stock.inv_io_bill WHERE bill_bn LIKE 'opening_in%'",
    "DELETE FROM oms_stock.inv_stock_ledger WHERE bm_id IN "
    "(SELECT id FROM oms_stock.inv_material WHERE material_bn LIKE 'IMP\\_%')",
    "DELETE FROM oms_stock.inv_material WHERE material_bn LIKE 'IMP\\_%'",
    "DELETE FROM oms_stock.inv_barcode WHERE barcode LIKE 'IMP\\_%'",
    "DELETE FROM oms_stock.inv_material_cat WHERE cat_name LIKE '导入%'",
    "DELETE FROM oms_stock.inv_brand WHERE brand_name LIKE '导入%'",
    "DELETE FROM oms_gateway.gw_shop WHERE shop_name LIKE '导入%'",
    "DELETE FROM oms_fulfill.ful_branch WHERE branch_name LIKE '导入%'",
    "DELETE FROM oms_org.org_node WHERE org_code LIKE 'WH\\_%' OR org_code LIKE 'SHOP\\_%'",
    "DELETE FROM oms_org.org_owner WHERE owner_code LIKE 'IMP\\_%'",
]


async def _cleanup() -> None:
    async with engine.begin() as conn:
        for sql in CLEANUP_SQL:
            await conn.execute(text(sql))


@pytest.fixture(autouse=True, scope="module")
async def isolate_import_data():
    await _cleanup()
    yield
    await _cleanup()


async def _owner_co_node_id() -> int:
    async with engine.begin() as conn:
        r = await conn.execute(
            text("SELECT id FROM oms_org.org_node "
                 "WHERE node_type='owner_co' AND is_deleted=0 LIMIT 1")
        )
        return int(r.scalar())


async def _make_owner(auth_client, code: str, name: str) -> dict:
    r = await auth_client.post("/api/v1/master/owners", json={
        "owner_code": code, "owner_name": name, "owner_type": "KC_BUYOUT",
        "org_node_id": await _owner_co_node_id(),
    })
    assert r.status_code == 200, r.text
    return r.json()["data"]


def _xlsx(headers: list[str], rows: list[list]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

MATERIAL_HEADERS = [
    "货号", "物料名称", "货主编码", "物料属性", "分类名称", "品牌名称",
    "SPU款号", "颜色", "尺码", "开票税率", "金蝶编码", "备注",
]
INITIAL_HEADERS = [
    "仓库编码", "货号", "货主编码", "实际库存", "锁定库存", "在途库存",
    "次品库存", "安全库存", "单位成本", "批次号", "效期",
]


# ============ 模板（FR-01）============

async def test_templates_listed(auth_client):
    """模板列表含商品与期初库存两套"""
    r = await auth_client.get("/api/v1/import/templates")
    assert r.status_code == 200, r.text
    codes = {t["template_code"] for t in r.json()["data"]}
    assert "material_import" in codes
    assert "initial_stock_import" in codes


async def test_template_file_download(auth_client):
    """模板文件可下载（含表头 + 字段说明 sheet）"""
    r = await auth_client.get("/api/v1/import/templates/material_import/file")
    assert r.status_code == 200, r.text
    assert len(r.content) > 1000
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(r.content))
    assert "导入数据" in wb.sheetnames
    assert "字段说明" in wb.sheetnames
    wb.close()


# ============ 商品导入（FR-02d + FR-03）============

async def test_material_import_success(auth_client):
    """全部合法 → 状态=2（成功）+ 物料落库"""
    owner = await _make_owner(auth_client, "IMP_OWNER_1", "导入货主1")
    content = _xlsx(MATERIAL_HEADERS, [
        ["IMP_M1", "导入物料甲", owner["owner_code"], "3", "", "", "SPU1", "红", "M", "13", "KD1", ""],
        ["IMP_M2", "导入物料乙", owner["owner_code"], "1", "", "", "", "", "", "13", "", ""],
    ])
    r = await auth_client.post(
        "/api/v1/import/tasks",
        files={"file": ("m.xlsx", content, _XLSX_MIME)},
        data={"template_code": "material_import", "creator": "tester"},
    )
    assert r.status_code == 200, r.text
    task_id = r.json()["data"]["id"]

    run = await auth_client.post(f"/api/v1/import/tasks/{task_id}/run")
    assert run.status_code == 200, run.text
    d = run.json()["data"]
    assert d["status"] == 2, d
    assert d["result"]["success_rows"] == 2
    assert d["result"]["fail_rows"] == 0

    mats = (await auth_client.get("/api/v1/material/materials", params={
        "keyword": "IMP_M"})).json()["data"]
    assert len(mats) == 2
    assert all(m["source"] == "import" for m in mats)


async def test_material_import_partial_failure(auth_client):
    """行级错误不阻塞其他行 → 状态=3（部分成功）+ 错误明细落库"""
    owner = await _make_owner(auth_client, "IMP_OWNER_2", "导入货主2")
    content = _xlsx(MATERIAL_HEADERS, [
        ["IMP_OK1", "正常行", owner["owner_code"], "3", "", "", "", "", "", "13", "", ""],
        ["IMP_BAD1", "货主不存在", "NO_SUCH_OWNER", "3", "", "", "", "", "", "13", "", ""],
        ["IMP_BAD2", "属性非法", owner["owner_code"], "9", "", "", "", "", "", "13", "", ""],
    ])
    r = await auth_client.post(
        "/api/v1/import/tasks",
        files={"file": ("m2.xlsx", content, _XLSX_MIME)},
        data={"template_code": "material_import", "creator": "tester"},
    )
    task_id = r.json()["data"]["id"]
    run = await auth_client.post(f"/api/v1/import/tasks/{task_id}/run")
    d = run.json()["data"]
    assert d["status"] == 3, d
    assert d["result"]["success_rows"] == 1
    assert d["result"]["fail_rows"] == 2

    errs = (await auth_client.get(f"/api/v1/import/tasks/{task_id}/errors")).json()["data"]
    codes = {e["error_code"] for e in errs}
    assert "FORMAT_ENUM" in codes      # 格式层：属性枚举非法
    assert "BIZ_OWNER_NOT_FOUND" in codes  # 业务层：货主不存在
    # 错误消息含定位信息（行号 + 值）
    assert any("第3行" in e["error_message"] for e in errs)
    # 错误聚类摘要
    clusters = d["result"]["cluster_summary"]
    assert clusters and clusters[0]["count"] >= 1


async def test_material_import_duplicate_bn_in_file(auth_client):
    """文件内货号重复 → 行错误（不整批失败）"""
    owner = await _make_owner(auth_client, "IMP_OWNER_3", "导入货主3")
    content = _xlsx(MATERIAL_HEADERS, [
        ["IMP_DUP", "行一", owner["owner_code"], "3", "", "", "", "", "", "13", "", ""],
        ["IMP_DUP", "行二", owner["owner_code"], "3", "", "", "", "", "", "13", "", ""],
    ])
    r = await auth_client.post(
        "/api/v1/import/tasks",
        files={"file": ("m3.xlsx", content, _XLSX_MIME)},
        data={"template_code": "material_import", "creator": "tester"},
    )
    task_id = r.json()["data"]["id"]
    d = (await auth_client.post(f"/api/v1/import/tasks/{task_id}/run")).json()["data"]
    assert d["result"]["success_rows"] == 1
    assert d["result"]["fail_rows"] == 1
    errs = (await auth_client.get(f"/api/v1/import/tasks/{task_id}/errors")).json()["data"]
    assert errs[0]["error_code"] == "BIZ_MATERIAL_BN_DUPLICATE_IN_FILE"


async def test_same_file_md5_warns(auth_client):
    """同 MD5 24h 内重复创建任务 → 409 提示（FR-03d）"""
    owner = await _make_owner(auth_client, "IMP_OWNER_4", "导入货主4")
    content = _xlsx(MATERIAL_HEADERS, [
        ["IMP_MD5", "MD5物料", owner["owner_code"], "3", "", "", "", "", "", "13", "", ""],
    ])
    files = {"file": ("m4.xlsx", content, _XLSX_MIME)}
    data = {"template_code": "material_import", "creator": "tester"}
    r1 = await auth_client.post("/api/v1/import/tasks", files=files, data=data)
    assert r1.status_code == 200, r1.text
    r2 = await auth_client.post("/api/v1/import/tasks", files=files, data=data)
    assert r2.status_code == 409
    assert "24" in r2.json()["message"]


# ============ 期初库存导入 + 期初建账（D04.04 §5.4）============

async def _setup_branch(auth_client, name: str, owner_id: int) -> dict:
    """建仓 + 配置可服务货主（M1b-0 关系）"""
    r = await auth_client.post("/api/v1/master/branches", json={
        "branch_name": name, "owner_id": None, "is_ctrl": 1, "countable_flag": 1,
    })
    assert r.status_code == 200, r.text
    branch = r.json()["data"]
    bind = await auth_client.post(f"/api/v1/master/branches/{branch['id']}/owners", json={
        "owner_ids": [owner_id],
    })
    assert bind.status_code == 200, bind.text
    return branch


async def test_initial_stock_import_creates_bill_and_ledger(auth_client):
    """期初导入 → 生成期初单据 + 明细 + 记账行 + 流水"""
    owner = await _make_owner(auth_client, "IMP_OWNER_5", "导入货主5")
    branch = await _setup_branch(auth_client, "导入测试仓1", owner["id"])
    mat = (await auth_client.post("/api/v1/material/materials", json={
        "material_bn": "IMP_INIT_1", "material_name": "期初物料", "owner_id": owner["id"],
    })).json()["data"]

    content = _xlsx(INITIAL_HEADERS, [
        [branch["branch_bn"], "IMP_INIT_1", owner["owner_code"], 100, 10, 5, 2, 20, 19.9, "B001", ""],
    ])
    r = await auth_client.post(
        "/api/v1/import/tasks",
        files={"file": ("init.xlsx", content, _XLSX_MIME)},
        data={"template_code": "initial_stock_import", "creator": "tester"},
    )
    assert r.status_code == 200, r.text
    task_id = r.json()["data"]["id"]
    d = (await auth_client.post(f"/api/v1/import/tasks/{task_id}/run")).json()["data"]
    assert d["status"] == 2, d
    assert d["result"]["success_rows"] == 1

    # 记账行（五口径）
    async with engine.begin() as conn:
        led = (await conn.execute(text(
            "SELECT store, freeze, arrive_store, defective, safe_store, unit_cost "
            "FROM oms_stock.inv_stock_ledger WHERE branch_id=:b AND bm_id=:m AND owner_id=:o"),
            {"b": branch["id"], "m": mat["id"], "o": owner["id"]},
        )).fetchone()
        assert led is not None, "期初记账行未生成"
        assert led[0] == 100
        assert led[1] == 10
        assert led[2] == 5
        assert led[3] == 2
        assert led[4] == 20
        assert float(led[5]) == 19.9

        # 期初单据（类型 opening_in，不参与结算）
        bill = (await conn.execute(text(
            "SELECT bill_type, direction, status, settle_status, source FROM oms_stock.inv_io_bill "
            "WHERE bill_bn LIKE 'opening_in%' LIMIT 1"))).fetchone()
        assert bill is not None
        assert bill[0] == "opening_in"
        assert bill[1] == 1
        assert bill[3] == 0  # 不参与结算（§5.4）

        # 变动流水（口径级 before/after）
        logs = (await conn.execute(text(
            "SELECT change_field, change_num, before_balance, after_balance, action, source "
            "FROM oms_stock.inv_stock_change_log WHERE bill_bn LIKE 'opening_in%'"))).fetchall()
        assert len(logs) == 4  # store/freeze/arrive/defective
        store_log = next(x for x in logs if x[0] == "store")
        assert store_log[1] == 100 and store_log[2] == 0 and store_log[3] == 100
        assert store_log[5] == "migration"


async def test_initial_stock_reject_branch_owner_not_allowed(auth_client):
    """仓未配置可服务货主 → 行错误（三元记账前提）"""
    owner = await _make_owner(auth_client, "IMP_OWNER_6", "导入货主6")
    other = await _make_owner(auth_client, "IMP_OWNER_7", "导入货主7")
    r = await auth_client.post("/api/v1/master/branches", json={
        "branch_name": "导入测试仓2", "owner_id": None, "is_ctrl": 1, "countable_flag": 1,
    })
    branch = r.json()["data"]
    # 只绑 other 货主，导入用 owner → 应被拒
    await auth_client.post(f"/api/v1/master/branches/{branch['id']}/owners", json={
        "owner_ids": [other["id"]],
    })
    await auth_client.post("/api/v1/material/materials", json={
        "material_bn": "IMP_INIT_2", "material_name": "期初物料2", "owner_id": owner["id"],
    })

    content = _xlsx(INITIAL_HEADERS, [
        [branch["branch_bn"], "IMP_INIT_2", owner["owner_code"], 50, 0, 0, 0, 0, 10, "", ""],
    ])
    task = await auth_client.post(
        "/api/v1/import/tasks",
        files={"file": ("init2.xlsx", content, _XLSX_MIME)},
        data={"template_code": "initial_stock_import", "creator": "tester"},
    )
    task_id = task.json()["data"]["id"]
    d = (await auth_client.post(f"/api/v1/import/tasks/{task_id}/run")).json()["data"]
    assert d["status"] == 4 or d["result"]["fail_rows"] == 1
    errs = (await auth_client.get(f"/api/v1/import/tasks/{task_id}/errors")).json()["data"]
    assert errs[0]["error_code"] == "BIZ_BRANCH_OWNER_NOT_ALLOWED"


async def test_initial_stock_reject_freeze_gt_store(auth_client):
    """锁定库存 > 实际库存 → 行错误"""
    owner = await _make_owner(auth_client, "IMP_OWNER_8", "导入货主8")
    branch = await _setup_branch(auth_client, "导入测试仓3", owner["id"])
    await auth_client.post("/api/v1/material/materials", json={
        "material_bn": "IMP_INIT_3", "material_name": "期初物料3", "owner_id": owner["id"],
    })
    content = _xlsx(INITIAL_HEADERS, [
        [branch["branch_bn"], "IMP_INIT_3", owner["owner_code"], 10, 30, 0, 0, 0, 10, "", ""],
    ])
    task = await auth_client.post(
        "/api/v1/import/tasks",
        files={"file": ("init3.xlsx", content, _XLSX_MIME)},
        data={"template_code": "initial_stock_import", "creator": "tester"},
    )
    task_id = task.json()["data"]["id"]
    (await auth_client.post(f"/api/v1/import/tasks/{task_id}/run"))
    errs = (await auth_client.get(f"/api/v1/import/tasks/{task_id}/errors")).json()["data"]
    assert errs[0]["error_code"] == "BIZ_FREEZE_GT_STORE"
